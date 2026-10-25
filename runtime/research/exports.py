import base64,hashlib,io,json,os,re,subprocess,tempfile,zipfile
from pathlib import Path
from xml.etree import ElementTree
from pypdf import PdfReader
from docx import Document
from citeproc_styles import get_style_filepath

STYLES={'apa':'apa','ieee':'ieee','acm':'association-for-computing-machinery','vancouver':'vancouver-nlm','chicago':'chicago-author-date','harvard':'harvard-cite-them-right'}

def escape(text):
    return re.sub(r'([\\`*_{}\[\]<>#$|~])',r'\\\1',str(text)).replace('\n',' ')

def equation(text):
    if len(text)>4000 or '^^' in text:raise ValueError('unsafe_equation')
    allowed=set('frac sqrt sum prod int lim left right cdot times div pm mp le ge neq approx infty alpha beta gamma delta epsilon theta lambda mu pi sigma phi omega partial nabla sin cos tan log ln exp begin end text mathrm mathbf mathit overline underline hat bar vec quad qquad cdots ldots dots forall exists in notin subset cup cap to rightarrow Leftrightarrow matrix pmatrix bmatrix aligned cases split equation align operatorname'.split())
    if any(c not in allowed for c in re.findall(r'\\([A-Za-z]+)',text)):raise ValueError('unsafe_equation')
    if any(env not in ['matrix','pmatrix','bmatrix','aligned','cases','split','equation','align'] for env in re.findall(r'\\(?:begin|end)\{([^}]+)\}',text)):raise ValueError('unsafe_equation_environment')
    return text.replace('$','')

def csl(source):
    authors=[]
    for a in source.get('authors',[]):
        if a.get('family'):authors.append({k:a[k] for k in ['given','family'] if a.get(k)})
        elif a.get('literal'):authors.append({'literal':a['literal']})
    v={'id':'s'+str(source['id']),'type':'article-journal','title':source['title'],'author':authors}
    for src,dst in [('doi','DOI'),('venue','container-title'),('url','URL')]:
        if source.get(src):v[dst]=source[src]
    if source.get('year'):v['issued']={'date-parts':[[int(source['year'])]]}
    return v

def export(payload):
    fmt=payload['format']
    if fmt not in ['docx','pdf','tex','bib','md']:raise ValueError('unsupported_format')
    settings=payload.get('settings',{});style=STYLES.get(settings.get('citation_style','apa'))
    if not style:raise ValueError('unsupported_csl')
    footnotes=[];cited=set();figures={int(f['id']):f for f in payload.get('figures',[])};source_ids={int(s['id']) for s in payload['sources']}
    with tempfile.TemporaryDirectory() as tmp:
        root=Path(tmp)
        def node(n):
            t=n.get('type');children=n.get('content',[])
            if t=='text':
                text=escape(n.get('text',''))
                for mark in n.get('marks',[]):
                    text={'bold':'**','italic':'*','strike':'~~','code':'`'}.get(mark['type'],'')+text+{'bold':'**','italic':'*','strike':'~~','code':'`'}.get(mark['type'],'')
                return text
            if t=='citation':
                sid=int(n['attrs']['sourceId'])
                if sid not in source_ids:raise ValueError('unresolved_citation')
                cited.add(sid);return ' [@s'+str(sid)+']'
            if t=='math':return '$'+equation(n['attrs']['latex'])+'$'
            if t=='footnote':footnotes.append(escape(n['attrs']['text']));return '[^f'+str(len(footnotes))+']'
            if t=='crossReference':return '[Section reference](#'+re.sub('[^a-zA-Z0-9_-]','',n['attrs']['target'])+')'
            if t=='figure':
                f=figures.get(int(n['attrs']['artifactId']))
                if not f:raise ValueError('missing_figure')
                raw=base64.b64decode(f['data'],validate=True)
                ext='png' if raw.startswith(b'\x89PNG\r\n\x1a\n') else 'jpg' if raw.startswith(b'\xff\xd8\xff') else None
                if not ext or len(raw)>4_000_000:raise ValueError('invalid_figure')
                path='figure-'+str(f['id'])+'.'+ext;(root/path).write_bytes(raw)
                url='data:image/'+('png' if ext=='png' else 'jpeg')+';base64,'+f['data'] if fmt=='md' else path
                return '\n\n!['+escape(n['attrs'].get('caption','Figure'))+']('+url+')\n\n'
            if t=='table':
                rows=[[ ''.join(node(c) for c in cell.get('content',[])).strip().replace('\n',' ') for cell in row.get('content',[])] for row in children]
                if not rows:return ''
                width=max(map(len,rows));rows=[row+['']*(width-len(row)) for row in rows]
                return '\n\n'+'| '+' | '.join(rows[0])+' |\n| '+' | '.join(['---']*width)+' |\n'+''.join('| '+' | '.join(row)+' |\n' for row in rows[1:])+'\n'
            text=''.join(node(c) for c in children)
            if t=='heading':return '\n\n'+'#'*max(1,min(6,int(n.get('attrs',{}).get('level',1))))+' '+text+' {#'+re.sub('[^a-zA-Z0-9_-]','',str(n.get('attrs',{}).get('id','section')))+'}\n\n'
            if t=='paragraph':return text+'\n\n'
            if t in ['bulletList','orderedList']:return '\n'+''.join(('- ' if t=='bulletList' else str(i+1)+'. ')+node(c).strip()+'\n' for i,c in enumerate(children))+'\n'
            if t=='blockquote':return '\n> '+text.replace('\n','\n> ')+'\n'
            if t=='codeBlock':return '\n~~~\n'+''.join(str(c.get('text','')) for c in children)+'\n~~~\n'
            if t=='hardBreak':return '  \n'
            if t=='horizontalRule':return '\n\n---\n\n'
            return text
        content=node(payload['document'])
        if payload.get('draft',True):content='**DRAFT — unresolved checks may remain. AI assistance requires human review.**\n\n'+content
        content+='\n'+''.join('[^f'+str(i+1)+']: '+text+'\n' for i,text in enumerate(footnotes))
        refs=[csl(s) for s in payload['sources'] if fmt=='bib' or int(s['id']) in cited]
        (root/'references.json').write_text(json.dumps(refs,ensure_ascii=False));(root/'manuscript.md').write_text(content)
        stylefile=get_style_filepath(style)
        def run(args):
            p=subprocess.run(['pandoc',*args],cwd=root,stdout=subprocess.PIPE,stderr=subprocess.PIPE,timeout=100)
            if p.returncode:raise ValueError('document_conversion_failed: '+p.stderr.decode(errors='replace')[-300:])
            return p.stdout
        common=['manuscript.md','--from=markdown+pipe_tables+footnotes+tex_math_dollars+citations-raw_html-raw_tex','--citeproc','--bibliography=references.json','--csl='+str(stylefile),'--resource-path='+str(root)]
        ext=fmt
        if fmt=='bib':
            raw=run(['references.json','-f','csljson','-t','bibtex']);(root/'out.bib').write_bytes(raw);run(['out.bib','-f','bibtex','-t','csljson'])
        elif fmt=='md':raw=run(common+['-t','markdown-citations-raw_html']);assert raw
        elif fmt=='docx':
            run(common+['-o','out.docx']);raw=(root/'out.docx').read_bytes();d=Document(io.BytesIO(raw))
            # Native bidi flags for Arabic paragraphs; scripts remain Unicode text.
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            for para in d.paragraphs:
                if re.search('[\u0600-\u06ff]',para.text):
                    bidi=OxmlElement('w:bidi');bidi.set(qn('w:val'),'1');para._p.get_or_add_pPr().append(bidi)
            bio=io.BytesIO();d.save(bio);raw=bio.getvalue();Document(io.BytesIO(raw))
            with zipfile.ZipFile(io.BytesIO(raw))as z:ElementTree.fromstring(z.read('word/document.xml'))
        else:
            font=int(settings.get('font_size',11));margin=int(settings.get('margin_mm',25));columns=int(settings.get('columns',1));journal=settings.get('journal','generic')
            if journal in ['ieee','acm']:columns=2;font=10;margin=19
            if not 9<=font<=14 or not 15<=margin<=35 or columns not in [1,2]:raise ValueError('invalid_layout')
            (root/'header.tex').write_text('\\usepackage{fontspec}\n')
            variables=['-V',f'fontsize={font}pt','-V',f'geometry:margin={margin}mm','-V','mainfont=DejaVu Serif']
            if columns==2:variables+=['-V','classoption=twocolumn']
            if fmt=='pdf':
                run(common+variables+['--pdf-engine=xelatex','-o','out.pdf']);raw=(root/'out.pdf').read_bytes();reader=PdfReader(io.BytesIO(raw));assert len(reader.pages)>0
            else:
                run(common+variables+['--standalone','-t','latex','-o','manuscript.tex']);tex=(root/'manuscript.tex').read_text()
                if '\\begin{document}' not in tex:raise ValueError('invalid_latex')
                bib=run(['references.json','-f','csljson','-t','bibtex']);(root/'references.bib').write_bytes(bib)
                (root/'README.txt').write_text('Compile manuscript.tex with XeLaTeX. Independent formatting preset; check target journal requirements. References are citeproc-formatted in the source. Figures are bundled. AI assistance and unresolved checks require review.')
                bio=io.BytesIO()
                with zipfile.ZipFile(bio,'w',zipfile.ZIP_DEFLATED)as z:
                    for f in root.iterdir():
                        if f.name in ['manuscript.tex','references.bib','README.txt']or f.name.startswith('figure-'):z.write(f,f.name)
                raw=bio.getvalue();ext='zip'
        if len(raw)>15_000_000:raise ValueError('export_size_limit')
        return {'data':base64.b64encode(raw).decode(),'sha256':hashlib.sha256(raw).hexdigest(),'validated':True,'extension':ext}
